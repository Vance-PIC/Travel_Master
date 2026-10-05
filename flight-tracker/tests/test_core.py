from __future__ import annotations

import sqlite3
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from src.db.database import (
    clear_all_price_snapshots,
    clear_all_tracking_data,
    count_price_snapshots,
    fetch_flight_instances,
    fetch_roundtrip_snapshots,
    get_connection,
    init_db,
    insert_roundtrip_snapshot,
    insert_snapshot,
    upsert_flight_instance,
)
from src.db.models import FlightInstance, FlightSnapshot, RoundTripSnapshot
from src.pipeline.candidate_selector import build_roundtrip_jobs
from src.pipeline.filter_engine import passes_strict_rule
from src.pipeline.runner import _build_scrapers, _crawl_tasks, _windows_to_run, run_pipeline
from src.pipeline.roundtrip_service import enrich_roundtrip_prices
from src.pipeline.schedule_builder import build_roundtrip_schedule_rows, build_schedule_rows
from src.scraper.skyscanner_crawler import SkyscannerCrawler


class FilterRuleTests(unittest.TestCase):
    def test_strict_boundaries(self):
        self.assertTrue(passes_strict_rule("2027-07-11", "OUTBOUND", "14:59"))
        self.assertFalse(passes_strict_rule("2027-07-11", "OUTBOUND", "15:00"))
        self.assertTrue(passes_strict_rule("2027-07-17", "INBOUND", "19:59"))
        self.assertFalse(passes_strict_rule("2027-07-18", "INBOUND", "14:00"))


class DatabaseTests(unittest.TestCase):
    def test_old_schema_migration_and_dedup(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            db = Path(tmpdir) / "test.sqlite3"
            raw = sqlite3.connect(db)
            raw.execute(
                "CREATE TABLE flight_price_snapshots ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT, flight_date TEXT NOT NULL, "
                "trip_type TEXT NOT NULL, origin_airport TEXT NOT NULL, dest_airport TEXT NOT NULL, "
                "airline_code TEXT, flight_number TEXT NOT NULL, departure_time TEXT NOT NULL, "
                "arrival_time TEXT NOT NULL, price_twd INTEGER, original_currency TEXT, "
                "original_price REAL, snapshot_hash TEXT UNIQUE, "
                "captured_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)"
            )
            raw.close()
            init_db(str(db))

            conn = get_connection(str(db))
            columns = {row["name"] for row in conn.execute("PRAGMA table_info(flight_price_snapshots)")}
            self.assertIn("last_checked_at", columns)
            snapshot = FlightSnapshot(
                flight_date="2027-07-11", trip_type="OUTBOUND", origin_airport="TPE",
                dest_airport="KIX", airline_code="BR", flight_number="BR108",
                departure_time="07:55", arrival_time="11:35", price_twd=7600,
            )
            self.assertTrue(insert_snapshot(conn, snapshot))
            self.assertFalse(insert_snapshot(conn, snapshot))
            count = conn.execute("SELECT COUNT(*) FROM flight_price_snapshots").fetchone()[0]
            self.assertEqual(1, count)
            conn.close()

    def test_roundtrip_snapshot_dedup_and_join(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            db = str(Path(tmpdir) / "roundtrip.sqlite3")
            init_db(db)
            conn = get_connection(db)
            outbound_id, _ = upsert_flight_instance(conn, FlightInstance(
                flight_date="2027-07-11", trip_type="OUTBOUND",
                origin_airport="TPE", dest_airport="KIX", airline_code="BR",
                flight_number="BR108", departure_time="07:55", arrival_time="11:35", stops=0,
            ))
            inbound_id, _ = upsert_flight_instance(conn, FlightInstance(
                flight_date="2027-07-17", trip_type="INBOUND",
                origin_airport="KIX", dest_airport="TPE", airline_code="BR",
                flight_number="BR109", departure_time="12:45", arrival_time="15:05", stops=0,
            ))
            fare = RoundTripSnapshot(
                outbound_flight_instance_id=outbound_id,
                inbound_flight_instance_id=inbound_id,
                window_key="W", price_twd=23100,
            )
            self.assertTrue(insert_roundtrip_snapshot(conn, fare))
            self.assertFalse(insert_roundtrip_snapshot(conn, fare))
            rows = fetch_roundtrip_snapshots(conn)
            self.assertEqual(1, len(rows))
            self.assertEqual("BR108", rows[0]["outbound_flight_number"])
            self.assertEqual(23100, rows[0]["price_twd"])
            self.assertEqual(
                {"single_leg": 0, "roundtrip": 1, "total": 1},
                count_price_snapshots(conn),
            )
            self.assertEqual(1, clear_all_price_snapshots(conn)["total"])
            conn.close()

    def test_clear_all_tracking_data_removes_prices_and_instances(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            db = str(Path(tmpdir) / "clear_all.sqlite3")
            init_db(db)
            conn = get_connection(db)
            outbound_id, _ = upsert_flight_instance(conn, FlightInstance(
                flight_date="2027-07-11", trip_type="OUTBOUND",
                origin_airport="TPE", dest_airport="KIX", airline_code="CI",
                flight_number="CI156", departure_time="08:15", arrival_time="12:00", stops=0,
            ))
            inbound_id, _ = upsert_flight_instance(conn, FlightInstance(
                flight_date="2027-07-17", trip_type="INBOUND",
                origin_airport="KIX", dest_airport="TPE", airline_code="CI",
                flight_number="CI157", departure_time="13:10", arrival_time="15:05", stops=0,
            ))
            insert_snapshot(conn, FlightSnapshot(
                flight_instance_id=outbound_id, flight_date="2027-07-11",
                trip_type="OUTBOUND", origin_airport="TPE", dest_airport="KIX",
                airline_code="CI", flight_number="CI156", departure_time="08:15",
                arrival_time="12:00", price_twd=32000,
            ))
            insert_roundtrip_snapshot(conn, RoundTripSnapshot(
                outbound_flight_instance_id=outbound_id,
                inbound_flight_instance_id=inbound_id,
                window_key="JULY", price_twd=64000,
            ))

            self.assertEqual(
                {"single_leg": 1, "roundtrip": 1, "flight_instances": 2, "total": 4},
                clear_all_tracking_data(conn),
            )
            self.assertEqual({"single_leg": 0, "roundtrip": 0, "total": 0}, count_price_snapshots(conn))
            self.assertEqual([], fetch_flight_instances(conn))
            conn.close()


class PipelineConfigurationTests(unittest.TestCase):
    def test_windows_and_sources_are_dynamic(self):
        settings = {"windows": {"CUSTOM": {}}, "origins": ["TPE", "TSA"]}
        self.assertEqual(["CUSTOM"], _windows_to_run("ALL", settings))
        self.assertEqual(
            ["google_flights", "skyscanner"],
            [s.source_name for s in _build_scrapers(["google_flights", "skyscanner"])],
        )

    def test_windows_follow_query_order_with_stable_ties(self):
        settings = {"windows": {
            "W_LATE": {"query_order": 9},
            "W_FIRST": {"query_order": 1},
            "W_FIRST_TIE": {"query_order": 1},
            "W_LEGACY": {},
        }}
        self.assertEqual(
            ["W_FIRST", "W_FIRST_TIE", "W_LEGACY", "W_LATE"],
            _windows_to_run("ALL", settings),
        )

    def test_crawl_tasks_include_all_origins(self):
        from src.config import load_settings
        settings = load_settings()
        window = next(iter(settings["windows"]))
        tasks = _crawl_tasks(["KIX"], [window])
        outbound_origins = {task[1] for task in tasks if task[4] == "OUTBOUND"}
        self.assertEqual(set(settings["origins"]), outbound_origins)

    def test_google_query_stays_machine_parseable(self):
        from urllib.parse import unquote
        from src.scraper.playwright_crawler import GoogleFlightsCrawler
        url = unquote(GoogleFlightsCrawler._build_search_url("TPE", "KIX", "2027-07-11", 4, True))
        self.assertIn("one way", url)
        self.assertNotIn("4 adults", url)
        self.assertNotIn("nonstop", url)

    def test_runner_skips_source_before_fetch_when_date_is_out_of_range(self):
        class FakeGoogleScraper:
            source_name = "google_flights"
            called = False

            def query_skip_reason(self, _flight_date):
                return "日期超出 Google Flights 可查詢範圍，已跳過"

            def fetch(self, *_args, **_kwargs):
                self.called = True
                raise AssertionError("超出範圍時不得呼叫 fetch")

        from src.config import load_settings
        window = next(iter(load_settings()["windows"]))
        scraper = FakeGoogleScraper()
        with tempfile.TemporaryDirectory() as tmpdir:
            test_db = str(Path(tmpdir) / "skip.sqlite3")
            with (
                patch("src.pipeline.runner._build_scrapers", return_value=[scraper]),
                patch("src.pipeline.runner._crawl_tasks", return_value=[
                    (window, "TPE", "KIX", "2099-01-01", "OUTBOUND")
                ]),
                patch("src.pipeline.runner.db_path", return_value=test_db),
            ):
                events = list(run_pipeline(
                    mode="crawl", destinations=["KIX"], window=window, stage="discover"
                ))
        self.assertFalse(scraper.called)
        self.assertTrue(any("已跳過" in event["message"] for event in events))

    def test_candidate_selection_requires_same_airline_and_caps_each_side(self):
        settings = {
            "origins": ["TPE"],
            "destinations": [{"iata": "KIX", "name_zh": "大阪"}],
            "windows": {"W": {
                "outbound_date": "2027-07-11", "inbound_dates": ["2027-07-17"],
            }},
            "airlines": {
                "BR": {"name_zh": "長榮航空", "is_lcc": False},
                "CI": {"name_zh": "中華航空", "is_lcc": False},
            },
            "scraper": {"roundtrip_candidate_limit_per_airline": 2},
            "strict_time_rules": {},
        }

        def instance(row_id, trip_type, airline, number, time):
            outbound = trip_type == "OUTBOUND"
            return {
                "id": row_id, "trip_type": trip_type, "airline_code": airline,
                "flight_number": number, "departure_time": time, "arrival_time": "12:00",
                "flight_date": "2027-07-11" if outbound else "2027-07-17",
                "origin_airport": "TPE" if outbound else "KIX",
                "dest_airport": "KIX" if outbound else "TPE", "stops": 0,
            }

        rows = [
            instance(1, "OUTBOUND", "BR", "BR1", "07:00"),
            instance(2, "OUTBOUND", "BR", "BR2", "08:00"),
            instance(3, "OUTBOUND", "BR", "BR3", "09:00"),
            instance(4, "INBOUND", "BR", "BR4", "12:00"),
            instance(5, "INBOUND", "BR", "BR5", "13:00"),
            instance(6, "OUTBOUND", "CI", "CI1", "10:00"),
        ]
        jobs = build_roundtrip_jobs(rows, ["KIX"], ["W"], settings)
        self.assertEqual(1, len(jobs))
        self.assertEqual("BR", jobs[0].airline_key)
        self.assertEqual(2, len(jobs[0].outbound))
        self.assertEqual(["BR5", "BR4"], [row["flight_number"] for row in jobs[0].inbound])

    def test_mock_enrichment_never_constructs_real_crawler(self):
        settings = {
            "origins": ["TPE"],
            "destinations": [{"iata": "KIX", "name_zh": "大阪"}],
            "windows": {"W": {
                "outbound_date": "2027-07-11", "inbound_dates": ["2027-07-17"],
            }},
            "airlines": {"BR": {"name_zh": "長榮航空", "is_lcc": False}},
            "scraper": {"roundtrip_candidate_limit_per_airline": 2},
            "strict_time_rules": {},
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            db = str(Path(tmpdir) / "mock.sqlite3")
            init_db(db)
            conn = get_connection(db)
            upsert_flight_instance(conn, FlightInstance(
                flight_date="2027-07-11", trip_type="OUTBOUND", origin_airport="TPE",
                dest_airport="KIX", airline_code="BR", flight_number="BR108",
                departure_time="07:55", arrival_time="11:35", stops=0,
            ))
            upsert_flight_instance(conn, FlightInstance(
                flight_date="2027-07-17", trip_type="INBOUND", origin_airport="KIX",
                dest_airport="TPE", airline_code="BR", flight_number="BR109",
                departure_time="12:45", arrival_time="15:05", stops=0,
            ))
            conn.close()

            def forbidden_crawler(**_kwargs):
                raise AssertionError("Mock 模式不得建立真實 crawler")

            updates = list(enrich_roundtrip_prices(
                db, ["KIX"], ["W"], 1, True, settings,
                mode="mock", crawler_factory=forbidden_crawler,
            ))
            conn = get_connection(db)
            fares = fetch_roundtrip_snapshots(conn)
            conn.close()
            self.assertTrue(updates)
            self.assertEqual(1, len(fares))
            self.assertEqual("mock", fares[0]["source"])

    def test_mock_all_pipeline_progress_never_goes_backwards(self):
        from src.config import load_settings

        window = next(iter(load_settings()["windows"]))
        with tempfile.TemporaryDirectory() as tmpdir:
            test_db = str(Path(tmpdir) / "pipeline.sqlite3")
            with patch("src.pipeline.runner.db_path", return_value=test_db):
                events = list(run_pipeline(
                    mode="mock", destinations=["KIX"], window=window,
                    stage="all", passenger_count=1, direct_only=True,
                ))
        percentages = [event["percent"] for event in events]
        self.assertEqual(percentages, sorted(percentages))
        self.assertEqual(100, percentages[-1])
        self.assertTrue(any("實際來回票價" in event["message"] for event in events))


class ScheduleTests(unittest.TestCase):
    def test_pairing_excludes_lcc_and_marks_strict_failure(self):
        settings = {
            "destinations": [{"iata": "KIX", "name_zh": "大阪"}],
            "windows": {"W": {"outbound_date": "2027-07-11", "inbound_dates": ["2027-07-18"]}},
        }
        base = {"airline_name": "長榮航空", "is_lcc": False, "price_twd": 5000}
        flights = [
            {**base, "dest_airport": "KIX", "origin_airport": "TPE", "trip_type": "OUTBOUND",
             "flight_date": "2027-07-11", "flight_number": "BR1", "departure_time": "10:00",
             "passes_strict_rule": True},
            {**base, "dest_airport": "TPE", "origin_airport": "KIX", "trip_type": "INBOUND",
             "flight_date": "2027-07-18", "flight_number": "BR2", "departure_time": "15:00",
             "passes_strict_rule": False},
        ]
        rows = build_schedule_rows(flights, settings)
        self.assertEqual(1, len(rows))
        self.assertIn("回程不符合", rows[0]["note"])

    def test_uses_latest_price_and_marks_one_overall_best(self):
        settings = {
            "destinations": [{"iata": "KIX", "name_zh": "大阪"}],
            "windows": {"W": {"group": "JULY", "outbound_date": "2027-07-11",
                               "inbound_dates": ["2027-07-18"]}},
        }

        def fare(instance_id, airline, number, trip_type, price, checked_at, passengers=4):
            outbound = trip_type == "OUTBOUND"
            return {
                "id": instance_id * 10 + price,
                "flight_instance_id": instance_id,
                "source": "google_flights",
                "cabin_class": "ECONOMY",
                "passenger_count": passengers,
                "last_checked_at": checked_at,
                "airline_name": airline,
                "is_lcc": False,
                "price_twd": price,
                "dest_airport": "KIX" if outbound else "TPE",
                "origin_airport": "TPE" if outbound else "KIX",
                "trip_type": trip_type,
                "flight_date": "2027-07-11" if outbound else "2027-07-18",
                "flight_number": number,
                "departure_time": "10:00" if outbound else "13:00",
                "passes_strict_rule": True,
            }

        flights = [
            # 同一航班、同一來源的歷史低價不得蓋過後來查到的現價。
            fare(1, "長榮航空", "BR1", "OUTBOUND", 1000, "2026-08-20 10:00:00"),
            fare(1, "長榮航空", "BR1", "OUTBOUND", 5000, "2026-08-25 10:00:00"),
            fare(2, "長榮航空", "BR2", "INBOUND", 4000, "2026-08-25 10:00:00"),
            fare(3, "中華航空", "CI1", "OUTBOUND", 4000, "2026-08-25 10:00:00"),
            fare(4, "中華航空", "CI2", "INBOUND", 4000, "2026-08-25 10:00:00"),
            # 不同旅客人數不應混入四人配對。
            fare(5, "長榮航空", "BR3", "OUTBOUND", 500, "2026-08-25 10:00:00", passengers=1),
        ]

        rows = build_schedule_rows(flights, settings, passenger_count=4, cabin_class="ECONOMY")
        self.assertEqual(2, len(rows))
        eva = next(row for row in rows if row["outbound_airline"] == "長榮航空")
        china = next(row for row in rows if row["outbound_airline"] == "中華航空")
        self.assertEqual(9000, eva["price_twd_roundtrip"])
        self.assertTrue(eva["is_airline_best"])
        self.assertFalse(eva["is_recommended"])
        self.assertTrue(china["is_airline_best"])
        self.assertTrue(china["is_recommended"])
        self.assertEqual(1, sum(row["is_recommended"] for row in rows))

    def test_formal_schedule_uses_actual_total_and_excludes_mixed_airlines(self):
        settings = {
            "destinations": [{"iata": "KIX", "name_zh": "大阪"}],
            "windows": {"W": {
                "group": "JULY", "outbound_date": "2027-07-11",
                "inbound_dates": ["2027-07-17"],
            }},
            "airlines": {
                "BR": {"name_zh": "長榮航空", "is_lcc": False},
                "CI": {"name_zh": "中華航空", "is_lcc": False},
            },
            "strict_time_rules": {},
        }

        def fare(row_id, inbound_airline, price):
            return {
                "id": row_id, "window_key": "W", "source": "google_flights",
                "cabin_class": "ECONOMY", "passenger_count": 1, "price_twd": price,
                "last_checked_at": "2026-08-26 10:00:00",
                "outbound_flight_instance_id": 1,
                "inbound_flight_instance_id": row_id + 1,
                "outbound_date": "2027-07-11", "inbound_date": "2027-07-17",
                "origin_airport": "TPE", "dest_airport": "KIX",
                "outbound_airline_code": "BR", "inbound_airline_code": inbound_airline,
                "outbound_flight_number": "BR108", "inbound_flight_number": "BR109",
                "outbound_departure_time": "07:55", "outbound_arrival_time": "11:35",
                "inbound_departure_time": "12:45", "inbound_arrival_time": "15:05",
            }

        rows = build_roundtrip_schedule_rows([
            fare(1, "BR", 23123), fare(2, "CI", 10000),
        ], settings)
        self.assertEqual(1, len(rows))
        self.assertEqual(23123, rows[0]["price_twd_roundtrip"])
        self.assertTrue(rows[0]["is_recommended"])


class SkyscannerParserTests(unittest.TestCase):
    def test_captcha_detection(self):
        self.assertTrue(
            SkyscannerCrawler._is_challenge_page(
                "https://www.skyscanner.com.tw/sttc/px/captcha-v2/index.html",
                "Are you a person or a robot?",
            )
        )

    def test_card_parser(self):
        parsed = SkyscannerCrawler._parse_card(
            "長榮航空\nBR 108\n07:55\n11:35\nNT$ 7,600",
            "TPE", "KIX", "2027-07-11", "OUTBOUND",
        )
        self.assertIsNotNone(parsed)
        self.assertEqual("BR108", parsed.flight_number)
        self.assertEqual(7600, parsed.original_price)


class ExporterTests(unittest.TestCase):
    def test_excel_contains_single_leg_and_formal_roundtrip_sheets(self):
        from openpyxl import load_workbook
        from src.exporter.excel_exporter import export_snapshots_to_excel

        leg = FlightSnapshot(
            flight_date="2027-07-11", trip_type="OUTBOUND",
            origin_airport="TPE", dest_airport="KIX", airline_code="BR",
            flight_number="BR108", departure_time="07:55", arrival_time="11:35",
            price_twd=7600,
        )
        roundtrip = {
            "window_key": "W", "outbound_date": "2027-07-11",
            "origin_airport": "TPE", "dest_airport": "KIX",
            "outbound_airline_code": "BR", "outbound_flight_number": "BR108",
            "outbound_departure_time": "07:55", "inbound_date": "2027-07-17",
            "inbound_airline_code": "BR", "inbound_flight_number": "BR109",
            "inbound_departure_time": "12:45", "passenger_count": 1,
            "cabin_class": "ECONOMY", "price_twd": 23100,
            "source": "google_flights",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            with patch(
                "src.exporter.excel_exporter.reports_dir", return_value=Path(tmpdir),
            ):
                output = export_snapshots_to_excel([leg], roundtrip_snapshots=[roundtrip])
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(["FlightSnapshots", "RoundTripPrices"], workbook.sheetnames)
            self.assertEqual(23100, workbook["RoundTripPrices"]["N2"].value)
            workbook.close()


class SettingsApiTests(unittest.TestCase):
    def test_settings_save_normalizes_query_order_to_integer(self):
        from src.api import server

        current = {
            "origins": [], "destinations": [], "windows": {}, "airlines": {},
            "strict_time_rules": {}, "database": {}, "reports": {},
        }
        request = server.SettingsUpdateRequest(
            origins=["TPE"], destinations=[],
            windows={"W": {"query_order": "2"}}, airlines={}, strict_time_rules={},
        )
        saved = []
        with (
            patch("src.api.server.load_settings", return_value=current),
            patch("src.api.server.save_settings", side_effect=saved.append),
        ):
            result = server.api_update_settings(request)
        self.assertEqual(2, result["windows"]["W"]["query_order"])
        self.assertEqual(2, saved[0]["windows"]["W"]["query_order"])

    def test_settings_rejects_non_positive_query_order(self):
        from fastapi import HTTPException
        from src.api import server

        request = server.SettingsUpdateRequest(
            origins=["TPE"], destinations=[], windows={"W": {"query_order": 0}},
            airlines={}, strict_time_rules={},
        )
        with patch("src.api.server.load_settings", return_value={}):
            with self.assertRaises(HTTPException) as raised:
                server.api_update_settings(request)
        self.assertEqual(422, raised.exception.status_code)


class SnapshotApiTests(unittest.TestCase):
    def test_roundtrip_api_keeps_four_semantic_times(self):
        from src.api import server

        with tempfile.TemporaryDirectory() as tmpdir:
            db = str(Path(tmpdir) / "api.sqlite3")
            init_db(db)
            conn = get_connection(db)
            outbound_id, _ = upsert_flight_instance(conn, FlightInstance(
                flight_date="2027-07-11", trip_type="OUTBOUND", origin_airport="TPE",
                dest_airport="PUS", airline_code="KE", flight_number="KE2086",
                departure_time="12:00", arrival_time="15:30", stops=0,
            ))
            inbound_id, _ = upsert_flight_instance(conn, FlightInstance(
                flight_date="2027-07-18", trip_type="INBOUND", origin_airport="PUS",
                dest_airport="TPE", airline_code="KE", flight_number="KE2085",
                departure_time="09:00", arrival_time="10:30", stops=0,
            ))
            insert_roundtrip_snapshot(conn, RoundTripSnapshot(
                outbound_flight_instance_id=outbound_id,
                inbound_flight_instance_id=inbound_id,
                window_key="W", price_twd=34680, passenger_count=4,
            ))
            conn.close()
            with patch("src.api.server.db_path", return_value=db):
                data = server.api_db_snapshots()
        row = next(row for row in data["snapshots"] if row["record_type"] == "ROUNDTRIP")
        self.assertEqual("15:30", row["outbound_arrival_time"])
        self.assertEqual("09:00", row["inbound_departure_time"])
        self.assertEqual("10:30", row["inbound_arrival_time"])
        self.assertNotIn("arrival_time", row)


class GoogleFlightsParserTests(unittest.TestCase):
    def test_itinerary_metadata_waits_for_attached_not_visible(self):
        from src.scraper.playwright_crawler import GoogleFlightsCrawler

        observed = {}

        class First:
            def wait_for(self, **kwargs):
                observed.update(kwargs)

        class Locator:
            first = First()

        class Page:
            def locator(self, selector):
                observed["selector"] = selector
                return Locator()

        crawler = GoogleFlightsCrawler(timeout_ms=4321)
        crawler._wait_for_itinerary_metadata(Page(), "[data-flight]")
        self.assertEqual("attached", observed["state"])
        self.assertEqual(4321, observed["timeout"])
        self.assertEqual("[data-flight]", observed["selector"])

    def test_local_date_guard_uses_configured_advance_limit(self):
        from src.scraper.playwright_crawler import GoogleFlightsCrawler

        crawler = GoogleFlightsCrawler(max_advance_days=330)
        today = date(2026, 8, 25)
        self.assertIsNone(crawler.query_skip_reason(str(today + timedelta(days=330)), today=today))
        reason = crawler.query_skip_reason(str(today + timedelta(days=331)), today=today)
        self.assertIn("已跳過", reason)

    def test_roundtrip_query_contains_both_dates(self):
        from urllib.parse import unquote
        from src.scraper.playwright_crawler import GoogleFlightsCrawler

        url = unquote(GoogleFlightsCrawler._build_roundtrip_search_url(
            "TPE", "KIX", "2027-07-11", "2027-07-17",
        ))
        self.assertIn("from 2027-07-11 to 2027-07-17 round trip", url)

    def test_hidden_itinerary_provides_real_flight_number(self):
        from src.scraper.playwright_crawler import GoogleFlightsCrawler

        class Card:
            def get_attribute(self, _):
                return "5,628 新台幣起 搭乘亞洲航空 X的直達航班 下午3:40 晚上7:30 總交通時間：2 小時 50 分鐘"

            def evaluate(self, _):
                return '<div data-travelimpactmodelwebsiteurl="https://example.test/?itinerary=TPE-KIX-D7-378-20270711">直達</div>'

        row = GoogleFlightsCrawler._parse_card(Card(), "TPE", "KIX", "2027-07-11", "OUTBOUND")
        self.assertIsNotNone(row)
        self.assertEqual("D7378", row.flight_number)
        self.assertEqual(170, row.duration_minutes)
        self.assertEqual(0, row.stops)

    def test_return_card_parser_reads_complete_roundtrip_total(self):
        from src.scraper.playwright_crawler import GoogleFlightsCrawler

        class Link:
            def get_attribute(self, _name):
                return ""

        class Card:
            def inner_text(self):
                return "長榮航空\nBR109\n$16,216\n來回票價"

            def locator(self, _selector):
                return type("First", (), {"first": Link()})()

        class Item:
            def get_attribute(self, _name):
                return (
                    "https://example.test/?itinerary=KIX-TPE-BR-109-20270717"
                )

            def locator(self, _selector):
                return Card()

        class Locator:
            def count(self):
                return 1

            def nth(self, _index):
                return Item()

        class Page:
            def locator(self, _selector):
                return Locator()

        quotes = GoogleFlightsCrawler._parse_return_quotes(
            Page(), "TPE", "KIX", "2027-07-17", "BR108", {"BR109"},
            "https://google.test/roundtrip",
        )
        self.assertEqual(1, len(quotes))
        self.assertEqual(16216, quotes[0].price_twd)
        self.assertEqual("BR109", quotes[0].inbound_flight_number)

    def test_roundtrip_card_fallback_matches_times_and_expanded_flight_number(self):
        from src.scraper.playwright_crawler import GoogleFlightsCrawler

        collapsed = "清晨6:30 – 上午10:10\n長榮航空\nTPE–KIX\n來回票價"
        expanded = collapsed + "\nBoeing 787-10\nBR\u00a0178"
        self.assertTrue(GoogleFlightsCrawler._card_matches_candidate_times(
            collapsed, "06:30", "10:10",
        ))
        self.assertTrue(GoogleFlightsCrawler._text_contains_flight_number(
            expanded, "BR178",
        ))
        self.assertFalse(GoogleFlightsCrawler._text_contains_flight_number(
            expanded, "BR108",
        ))


if __name__ == "__main__":
    unittest.main()
