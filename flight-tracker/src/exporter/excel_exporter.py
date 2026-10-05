"""openpyxl 美化匯出 .xlsx 報表模組。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import reports_dir
from src.db.models import FlightSnapshot

_HEADERS = [
    ("flight_date", "飛行日期"),
    ("trip_type", "行程類型"),
    ("origin_airport", "出發機場"),
    ("dest_airport", "抵達機場"),
    ("airline_code", "航空公司"),
    ("flight_number", "航班號碼"),
    ("departure_time", "出發時刻"),
    ("arrival_time", "抵達時刻"),
    ("price_twd", "票價(TWD)"),
    ("original_currency", "原始幣別"),
    ("original_price", "原始價格"),
    ("captured_at", "抓取時間"),
    ("last_checked_at", "最後查詢時間"),
]

_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")

_ROUNDTRIP_HEADERS = [
    ("window_key", "時程"),
    ("outbound_date", "去程日期"),
    ("origin_airport", "出發機場"),
    ("dest_airport", "目的地"),
    ("outbound_airline_code", "去程航空"),
    ("outbound_flight_number", "去程航班"),
    ("outbound_departure_time", "去程出發"),
    ("inbound_date", "回程日期"),
    ("inbound_airline_code", "回程航空"),
    ("inbound_flight_number", "回程航班"),
    ("inbound_departure_time", "回程出發"),
    ("passenger_count", "旅客人數"),
    ("cabin_class", "艙等"),
    ("price_twd", "正式來回總價(TWD)"),
    ("source", "來源"),
    ("captured_at", "抓取時間"),
    ("last_checked_at", "最後查詢時間"),
]


def _write_dict_sheet(ws, rows: list[dict], headers: list[tuple[str, str]]) -> None:
    for col_idx, (_, label) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (field, _) in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(field))
    for col_idx, (field, label) in enumerate(headers, start=1):
        values = [len(str(row.get(field) or "")) for row in rows]
        ws.column_dimensions[get_column_letter(col_idx)].width = max([len(label), *values]) + 4
    ws.freeze_panes = "A2"


def export_snapshots_to_excel(
    snapshots: list[FlightSnapshot], feature: str = "flight_report",
    roundtrip_snapshots: list[dict] | None = None,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "FlightSnapshots"

    for col_idx, (_, label) in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    sorted_snapshots = sorted(snapshots, key=lambda s: s.price_twd or 0)
    for row_idx, snap in enumerate(sorted_snapshots, start=2):
        data = snap.to_dict()
        for col_idx, (field, _) in enumerate(_HEADERS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=data.get(field))

    for col_idx, (field, label) in enumerate(_HEADERS, start=1):
        max_len = max([len(label)] + [len(str(getattr(s, field, "") or "")) for s in sorted_snapshots] or [len(label)])
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    ws.freeze_panes = "A2"

    if roundtrip_snapshots is not None:
        roundtrip_ws = wb.create_sheet("RoundTripPrices")
        _write_dict_sheet(
            roundtrip_ws,
            sorted(roundtrip_snapshots, key=lambda row: row.get("price_twd") or 0),
            _ROUNDTRIP_HEADERS,
        )

    output_dir = reports_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"{feature}_{timestamp}.xlsx"
    wb.save(output_path)
    return output_path
